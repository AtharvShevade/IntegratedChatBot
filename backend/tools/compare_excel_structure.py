# backend/tools/compare_excel_structure.py — compares the structure (sheet
# names + column headers per sheet) of a user-uploaded Excel file against the
# original sample/template for that return, and reports what's missing.

from __future__ import annotations

import openpyxl


def _sheet_headers(ws) -> list[str]:
    """First non-empty row is treated as the header row."""
    for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
        cells = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
        if cells:
            return cells
    return []


def compare_excel_structure(sample_path: str, user_path: str) -> dict:
    sample_wb = openpyxl.load_workbook(sample_path, read_only=True, data_only=True)
    user_wb = openpyxl.load_workbook(user_path, read_only=True, data_only=True)

    sample_sheets = sample_wb.sheetnames
    user_sheets = user_wb.sheetnames

    missing_sheets = [s for s in sample_sheets if s not in user_sheets]
    extra_sheets = [s for s in user_sheets if s not in sample_sheets]

    sheet_reports = []
    for sheet_name in sample_sheets:
        if sheet_name in missing_sheets:
            continue
        sample_headers = _sheet_headers(sample_wb[sheet_name])
        user_headers = _sheet_headers(user_wb[sheet_name])

        missing_cols = [c for c in sample_headers if c not in user_headers]
        extra_cols = [c for c in user_headers if c not in sample_headers]

        if missing_cols or extra_cols:
            sheet_reports.append({
                "sheet": sheet_name,
                "missing_columns": missing_cols,
                "extra_columns": extra_cols,
            })

    sample_wb.close()
    user_wb.close()

    is_same = not missing_sheets and not extra_sheets and not sheet_reports

    return {
        "is_same_structure": is_same,
        "missing_sheets": missing_sheets,
        "extra_sheets": extra_sheets,
        "sheet_differences": sheet_reports,
    }


def format_report(result: dict) -> str:
    if result["is_same_structure"]:
        return "Structure matches: all sheets and columns are the same."

    lines = ["Structure mismatch found:"]
    for s in result["missing_sheets"]:
        lines.append(f"- Missing sheet: '{s}'")
    for s in result["extra_sheets"]:
        lines.append(f"- Extra sheet not in sample: '{s}'")
    for diff in result["sheet_differences"]:
        for c in diff["missing_columns"]:
            lines.append(f"- Sheet '{diff['sheet']}': missing column '{c}'")
        for c in diff["extra_columns"]:
            lines.append(f"- Sheet '{diff['sheet']}': extra column '{c}' not in sample")
    return "\n".join(lines)


if __name__ == "__main__":
    import sys

    sample_path, user_path = sys.argv[1], sys.argv[2]
    result = compare_excel_structure(sample_path, user_path)
    print(format_report(result))
